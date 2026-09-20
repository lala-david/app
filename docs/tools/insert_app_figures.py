"""지원서 사본에 앱 화면 그림과 유스케이스 다이어그램을 넣는다. 원본은 건드리지 않는다."""
import copy
import io
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
ASSETS = ROOT / "docs" / "assets"
FIGURE_WIDTH = Cm(14.0)
DIAGRAM_WIDTH = Cm(13.0)
MAX_PIXELS = 2200
CAPTION_COLOR = RGBColor(0x5B, 0x6B, 0x66)

# (본문 표 번호, 이 글로 시작하는 문단 뒤, 그림, 캡션, 폭)
PLACEMENTS = [
    (33, "이에 사운즈펀 브릿지는", "application/fig-child.png",
     "<그림 1. SoundsFun Bridge 앱 시제품(1주차 ‘Colors’)의 아동 화면>", FIGURE_WIDTH),
    (39, "기존 무료 콘텐츠는", "application/fig-journey.png",
     "<그림 2. 48주 소리 여행과 이번 주 소리 활동>", FIGURE_WIDTH),
    (39, " 접근성", "application/fig-parent.png",
     "<그림 3. 보호자 화면: 학습 기록, 어려워한 말과 대화 가이드, 휴대폰 알림>", FIGURE_WIDTH),
    (39, None, "uml/usecase.png",
     "<그림 4. MVP 유스케이스 다이어그램>", DIAGRAM_WIDTH),
]


def shrunk(path: Path) -> io.BytesIO:
    source = Image.open(path).convert("RGBA")
    image = Image.new("RGB", source.size, "white")
    image.paste(source, mask=source.getchannel("A"))
    if image.width > MAX_PIXELS:
        image = image.resize((MAX_PIXELS, round(image.height * MAX_PIXELS / image.width)), Image.LANCZOS)
    out = io.BytesIO()
    image.save(out, "PNG", optimize=True)
    out.seek(0)
    return out


def paragraph_after(anchor: Paragraph) -> Paragraph:
    element = copy.deepcopy(anchor._p)
    for child in list(element):
        if child.tag != qn("w:pPr"):
            element.remove(child)
    anchor._p.addnext(element)
    paragraph = Paragraph(element, anchor._parent)
    paragraph.style = "Normal"
    fmt = paragraph.paragraph_format
    fmt.left_indent = fmt.first_line_indent = Cm(0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return paragraph


def add_figure(anchor: Paragraph, image: Path, caption: str, width) -> Paragraph:
    picture = paragraph_after(anchor)
    picture.paragraph_format.space_before = Pt(10)
    picture.paragraph_format.space_after = Pt(4)
    picture.paragraph_format.keep_with_next = True
    picture.add_run().add_picture(shrunk(image), width=width)

    text = paragraph_after(picture)
    text.paragraph_format.space_after = Pt(12)
    run = text.add_run(caption)
    run.font.size = Pt(9)
    run.font.color.rgb = CAPTION_COLOR
    return text


def main(source: str, target: str) -> None:
    document = docx.Document(source)
    body = list(document.element.body)
    last = None
    for table_index, starts_with, image, caption, width in PLACEMENTS:
        cell = Table(body[table_index], document).rows[0].cells[0]
        if starts_with is None:
            anchor = last
        else:
            matches = [p for p in cell.paragraphs if p.text.startswith(starts_with)]
            if len(matches) != 1:
                raise SystemExit(f"문단을 찾지 못했습니다: {starts_with!r} ({len(matches)}개)")
            anchor = matches[0]
        last = add_figure(anchor, ASSETS / image, caption, width)
    document.save(target)
    print("saved", target)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
