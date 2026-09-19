"""
docs/src/*.md  →  docs/html/*.html  +  docs/docx/*.docx

사용법:  python docs/build.py            (전체)
         python docs/build.py 04         (번호로 시작하는 파일만)

Markdown 한 벌이 원본이고, HTML과 DOCX는 여기서 생성됩니다. 출력 파일은 직접 고치지 마세요.
"""
from __future__ import annotations

import base64
import io
import re
import sys
import urllib.request
import zlib
from dataclasses import dataclass, field
from pathlib import Path

import markdown
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

ROOT = Path(__file__).parent
SRC, HTML_DIR, DOCX_DIR = ROOT / "src", ROOT / "html", ROOT / "docx"
MERMAID_INK = "https://mermaid.ink/img/"

# 다이어그램 색은 앱 디자인 토큰(노랑·청록·코랄·보라)에서 가져온다
MERMAID_THEME = {
    "theme": "base",
    "themeVariables": {
        "fontFamily": "Pretendard, Malgun Gothic, sans-serif", "fontSize": "15px",
        "primaryColor": "#FFF2A8", "primaryBorderColor": "#E0A800", "primaryTextColor": "#2B2A33",
        "secondaryColor": "#BFEFE7", "secondaryBorderColor": "#20AFA3", "tertiaryColor": "#FFF9E6", "tertiaryBorderColor": "#F4BC16",
        "lineColor": "#6B6880", "textColor": "#2B2A33", "mainBkg": "#FFF2A8", "nodeBorder": "#E0A800",
        "clusterBkg": "#FFFBEF", "clusterBorder": "#F4BC16", "edgeLabelBackground": "#FFFFFF",
        "actorBkg": "#BFEFE7", "actorBorder": "#20AFA3", "actorTextColor": "#16423E", "actorLineColor": "#B9B6C9",
        "signalColor": "#4A4860", "signalTextColor": "#2B2A33", "labelBoxBkgColor": "#FFD0D7", "labelBoxBorderColor": "#F34E67", "labelTextColor": "#5A1320",
        "loopTextColor": "#5A1320", "noteBkgColor": "#EBC5FA", "noteBorderColor": "#B64FE0", "noteTextColor": "#3B1450",
        "activationBkgColor": "#FFE472", "activationBorderColor": "#E0A800", "sequenceNumberColor": "#FFFFFF",
        "classText": "#2B2A33", "attributeBackgroundColorEven": "#FFFFFF", "attributeBackgroundColorOdd": "#FFFBEF",
    },
}

# ───────────────────────────── HTML ─────────────────────────────

HTML_TEMPLATE = """<!doctype html>
<html lang="ko">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=IBM+Plex+Sans+KR:wght@400;500;600&family=IBM+Plex+Mono:wght@400;500&family=Jua&display=swap">
<style>{css}</style>
</head>
<body>
<nav class="docnav">{nav}</nav>
<main>
<div class="docmeta">SoundsFun Bridge · MVP 문서 v2 · {file_no}</div>
{body}
</main>
<script type="module">
  import mermaid from "https://cdn.jsdelivr.net/npm/mermaid@11/dist/mermaid.esm.min.mjs";
  mermaid.initialize({{ startOnLoad: true, theme: "base" }});
</script>
</body>
</html>
"""

CSS = """
:root{--bg:#F3F6F4;--sf:#fff;--sf2:#E9EFEB;--ink:#1B2B25;--mut:#5E6E67;--ln:#D5DED9;--ac:#1F7A5C;--ac2:#DDEFE6;--co:#E8603F}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--ink);font:15px/1.7 "IBM Plex Sans KR","Apple SD Gothic Neo","Malgun Gothic",sans-serif}
main{max-width:1000px;margin:0 auto;padding:24px 20px 80px}
.docnav{position:sticky;top:0;z-index:5;background:var(--bg);border-bottom:1px solid var(--ln);padding:0 20px}
.docnav ul{list-style:none;margin:0 auto;padding:0;max-width:1000px;display:flex;gap:2px;overflow-x:auto;scrollbar-width:none}
.docnav a{display:block;padding:10px 9px;font-size:.82rem;color:var(--mut);text-decoration:none;white-space:nowrap;border-bottom:2px solid transparent}
.docnav a.cur,.docnav a:hover{color:var(--ac);border-color:var(--ac)}
.docmeta{font-size:.76rem;letter-spacing:.1em;text-transform:uppercase;color:var(--ac);font-weight:600;margin-top:18px}
h1,h2,h3,h4{font-family:"Jua","IBM Plex Sans KR",sans-serif;font-weight:400;letter-spacing:-.01em;text-wrap:balance;margin:0}
h1{font-size:2.3rem;line-height:1.15;margin:6px 0 10px}
h2{font-size:1.6rem;margin-top:48px;padding-top:16px;border-top:1px solid var(--ln)}
h3{font-size:1.2rem;margin-top:28px;color:#145A43}
h4{font-size:1rem;margin-top:18px;font-family:inherit;font-weight:600}
p{margin:10px 0;max-width:72ch}
a{color:var(--ac)}
ul,ol{margin:8px 0;padding-left:22px;max-width:76ch}li{margin:3px 0}
code{font-family:"IBM Plex Mono",ui-monospace,monospace;font-size:.86em;background:var(--sf2);padding:1px 5px;border-radius:4px}
pre{background:var(--sf);border:1px solid var(--ln);border-radius:10px;padding:14px 16px;overflow-x:auto;font-size:.82rem;line-height:1.55}
pre code{background:none;padding:0}
pre.mermaid{text-align:center}
.diagram{max-width:none;text-align:center;background:var(--sf);border:1px solid var(--ln);border-radius:12px;padding:12px}
.diagram img{max-width:100%;height:auto}
figure.fig{margin:16px 0;text-align:center;background:var(--sf);border:1px solid var(--ln);border-radius:12px;padding:12px}
figure.fig img{max-width:100%;height:auto;border-radius:8px}
figure.screens{display:inline-block;width:230px;margin:6px;vertical-align:top}
figcaption{font-size:.8rem;color:var(--mut);margin-top:6px}
table{border-collapse:collapse;width:100%;font-size:.88rem;background:var(--sf);margin:14px 0}
.tw{overflow-x:auto;border:1px solid var(--ln);border-radius:10px;margin:14px 0}
.tw table{margin:0}
th,td{text-align:left;vertical-align:top;padding:9px 12px;border-bottom:1px solid var(--ln)}
th{background:var(--sf2);font-weight:600;font-size:.8rem;white-space:nowrap}
tr:last-child td{border-bottom:0}
blockquote{margin:14px 0;padding:10px 16px;border-left:3px solid var(--co);background:var(--sf);border-radius:0 10px 10px 0}
blockquote p{margin:0}
hr{border:0;border-top:1px solid var(--ln);margin:32px 0}
strong{font-weight:600}
@media print{.docnav{display:none}body{background:#fff}main{padding:0}}
"""


def md_to_html(text: str) -> str:
    def diagram(m):
        png = diagram_file(m.group(1))
        if png:
            return f'<p class="diagram"><img src="../assets/diagrams/{png.name}" alt="다이어그램"></p>'
        return f'<pre class="mermaid">\n{m.group(1)}</pre>'

    text = re.sub(r"```mermaid\n(.*?)```", diagram, text, flags=re.S)
    html = markdown.markdown(text, extensions=["tables", "fenced_code", "sane_lists", "toc"])
    html = re.sub(r'<p><img alt="([^"]*)" src="([^"]+)" ?/?></p>', lambda m: f'<figure class="fig {Path(m.group(2)).parent.name}"><img src="{m.group(2)}" alt="{m.group(1)}"><figcaption>{m.group(1)}</figcaption></figure>', html)
    return re.sub(r"(<table>.*?</table>)", r'<div class="tw">\1</div>', html, flags=re.S)


# ───────────────────────────── DOCX ─────────────────────────────

@dataclass
class Block:
    kind: str                       # h1 h2 h3 h4 p ul ol table code mermaid quote
    text: str = ""
    items: list = field(default_factory=list)
    rows: list = field(default_factory=list)
    level: int = 0


def parse_blocks(text: str) -> list[Block]:
    lines, blocks, i = text.splitlines(), [], 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            lang = line[3:].strip()
            j = i + 1
            while j < len(lines) and not lines[j].startswith("```"):
                j += 1
            blocks.append(Block("mermaid" if lang == "mermaid" else "code", "\n".join(lines[i + 1:j])))
            i = j + 1
            continue
        im = re.match(r"^!\[(.*?)\]\((.+?)\)\s*$", line)
        if im:
            blocks.append(Block("image", im.group(2), items=[im.group(1)]))
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            blocks.append(Block(f"h{len(m.group(1))}", m.group(2).strip()))
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(Block("table", rows=rows))
            continue
        if line.startswith(">"):
            quote = []
            while i < len(lines) and lines[i].startswith(">"):
                quote.append(lines[i][1:].strip())
                i += 1
            blocks.append(Block("quote", " ".join(quote)))
            continue
        lm = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", line)
        if lm:
            kind = "ol" if lm.group(2)[0].isdigit() else "ul"
            items = []
            while i < len(lines):
                lm = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", lines[i])
                if not lm:
                    break
                items.append((len(lm.group(1)) // 2, lm.group(3).strip()))
                i += 1
            blocks.append(Block(kind, items=items))
            continue
        if line.strip() == "" or line.strip() == "---":
            i += 1
            continue
        para = []
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,4}\s|\||>|```|\s*([-*]|\d+\.)\s)", lines[i]):
            para.append(lines[i].strip())
            i += 1
        blocks.append(Block("p", " ".join(para)))
    return blocks


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_inline(paragraph, text: str, size: int | None = None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            run = paragraph.add_run(part[2:-2]); run.bold = True
        elif part.startswith("`"):
            run = paragraph.add_run(part[1:-1]); run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        else:
            run = paragraph.add_run(part)
        if size:
            run.font.size = Pt(size)


def shade(cell, hex_color: str):
    tc = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color)
    tc.append(shd)


DIAGRAM_DIR = ROOT / "assets" / "diagrams"
IMAGE_WIDTHS = {"screens": 6.2, "uml": 16.5}  # 폴더별 DOCX 그림 폭(cm)


def diagram_file(code: str) -> Path | None:
    """mermaid 코드를 PNG로 렌더링해 assets/diagrams에 캐시한다 (내용 해시로 파일명)."""
    import hashlib
    DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)
    path = DIAGRAM_DIR / f"{hashlib.sha1((code + "|v2").encode()).hexdigest()[:12]}.png"
    if path.exists():
        return path
    state = __import__("json").dumps({"code": code, "mermaid": MERMAID_THEME}, ensure_ascii=False)
    payload = base64.urlsafe_b64encode(zlib.compress(state.encode(), 9)).decode()
    try:
        req = urllib.request.Request(MERMAID_INK + "pako:" + payload + "?type=png&bgColor=FFFFFF&width=2000", headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=40) as r:
            path.write_bytes(r.read())
        return path
    except Exception as e:
        print("  diagram render failed:", e)
        return None


def mermaid_png(code: str) -> bytes | None:
    path = diagram_file(code)
    return path.read_bytes() if path else None


def set_base_font(doc: Document):
    style = doc.styles["Normal"]
    style.font.name = "Malgun Gothic"; style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    for name, size, color in (("Title", 24, "1F7A5C"), ("Heading 1", 17, "145A43"), ("Heading 2", 13.5, "1B2B25"), ("Heading 3", 11.5, "1F7A5C")):
        st = doc.styles[name]
        st.font.name = "Malgun Gothic"; st.font.size = Pt(size); st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(2)


def blocks_to_docx(blocks: list[Block], out: Path, title: str):
    doc = Document()
    set_base_font(doc)
    doc.add_paragraph("SoundsFun Bridge · MVP 문서 v2").runs[0].font.color.rgb = RGBColor.from_string("1F7A5C")
    for b in blocks:
        if b.kind == "h1":
            doc.add_heading(b.text, level=0)
        elif b.kind in ("h2", "h3", "h4"):
            doc.add_heading(b.text, level=int(b.kind[1]) - 1)
        elif b.kind == "p":
            add_inline(doc.add_paragraph(), b.text)
        elif b.kind == "quote":
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.6)
            add_inline(p, b.text)
            for r in p.runs:
                r.italic = True
        elif b.kind in ("ul", "ol"):
            style = "List Bullet" if b.kind == "ul" else "List Number"
            for depth, item in b.items:
                p = doc.add_paragraph(style=style if depth == 0 else f"{style} 2")
                add_inline(p, item)
        elif b.kind == "code":
            p = doc.add_paragraph()
            run = p.add_run(b.text); run.font.name = "Consolas"; run.font.size = Pt(8.5)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            p.paragraph_format.left_indent = Cm(0.4)
        elif b.kind == "mermaid":
            png = mermaid_png(b.text)
            if png:
                doc.add_picture(io.BytesIO(png), width=Cm(16))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p = doc.add_paragraph(); run = p.add_run("[다이어그램 - HTML 버전에서 확인]\n" + b.text)
                run.font.name = "Consolas"; run.font.size = Pt(8.5)
        elif b.kind == "image":
            path = (SRC / b.text).resolve()
            if path.exists():
                doc.add_picture(str(path), width=Cm(float(IMAGE_WIDTHS.get(path.parent.name, 16))))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                if b.items and b.items[0]:
                    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = cap.add_run(b.items[0]); run.font.size = Pt(9); run.font.color.rgb = RGBColor.from_string("5E6E67")
        elif b.kind == "table" and b.rows:
            cols = max(len(r) for r in b.rows)
            t = doc.add_table(rows=len(b.rows), cols=cols); t.style = "Table Grid"
            for ri, row in enumerate(b.rows):
                for ci in range(cols):
                    cell = t.cell(ri, ci)
                    cell.text = ""
                    add_inline(cell.paragraphs[0], row[ci] if ci < len(row) else "", size=9)
                    if ri == 0:
                        shade(cell, "E9EFEB")
                        for r in cell.paragraphs[0].runs:
                            r.bold = True
            doc.add_paragraph()
    doc.core_properties.title = title
    doc.save(out)


# ───────────────────────────── build ─────────────────────────────

def title_of(text: str, fallback: str) -> str:
    m = re.search(r"^#\s+(.*)", text, flags=re.M)
    return m.group(1).strip() if m else fallback


def build(filter_prefix: str | None = None):
    sources = sorted(SRC.glob("*.md"))
    nav_items = [(s.stem, title_of(s.read_text(encoding="utf-8"), s.stem)) for s in sources]
    for src in sources:
        if filter_prefix and not src.stem.startswith(filter_prefix):
            continue
        text = src.read_text(encoding="utf-8")
        title = title_of(text, src.stem)
        nav = "<ul>" + "".join(
            f'<li><a class="{"cur" if stem == src.stem else ""}" href="{stem}.html">{stem[:2]} {t}</a></li>' for stem, t in nav_items
        ) + "</ul>"
        html = HTML_TEMPLATE.format(title=title, css=CSS, nav=nav, body=md_to_html(text), file_no=src.stem[:2])
        (HTML_DIR / f"{src.stem}.html").write_text(html, encoding="utf-8")
        blocks_to_docx(parse_blocks(text), DOCX_DIR / f"{src.stem}.docx", title)
        print(f"built  {src.stem}")


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else None)
